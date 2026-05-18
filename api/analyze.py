from fastapi import FastAPI, UploadFile, File
from fastapi.responses import JSONResponse
from io import BytesIO
from openpyxl import load_workbook

app = FastAPI()

@app.post("/api/analyze")
async def analyze(file: UploadFile = File(...)):
    contents = await file.read()
    try:
        buf = BytesIO(contents)
        wb = load_workbook(buf, read_only=True, data_only=True)
        sheets = wb.sheetnames
        
        result = {}
        for sheet in sheets:
            ws = wb[sheet]
            headers = []
            for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
                for cell in row:
                    if cell is not None and str(cell).strip():
                        headers.append(str(cell).strip())
            result[sheet] = list(dict.fromkeys(headers))
            
        return JSONResponse({"sheets": result})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=400)
